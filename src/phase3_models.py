"""Phase 3 - fit all forecasting models, build inference tables, diagnostics and the comparison summary.

Split (chronological, no shuffling):
  train      2023-01-29 .. 2023-12-31   (December surge included in estimation, with a dummy where possible)
  validation 2024-01-01 .. 2024-01-31   (multi-step forecast from end of train; used for ranking)
  test       2024-02-01 .. 2024-02-28   (models refit on train+val with the same procedure; forecasts saved for
                                          Phase 4 and NOT used for any Phase 3 choice)
Outputs
  outputs/excel/model_inference_tables.xlsx   one sheet per model + model_comparison_summary + split_definition
  outputs/data/phase3_forecasts_val.csv, phase3_forecasts_test.csv
  outputs/figures/phase3_*.png
"""
import time
import warnings

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font

from . import models as mdl
from .config import OUT_D, OUT_FIG, OUT_X, TARGET
from .data_access import load_daily
from .diagnostics import accuracy, residual_diagnostics, residual_plot
from .eda import C, _fmt_dates, _thousands

warnings.filterwarnings("ignore")

SPLIT = dict(train=("2023-01-29", "2023-12-31"), val=("2024-01-01", "2024-01-31"), test=("2024-02-01", "2024-02-28"))
INF_COLS = ["model_name", "parameter/order", "coefficient_or_component", "estimate", "std_error",
            "p_value_or_significance", "AIC", "BIC", "MAE", "RMSE", "MAPE", "notes"]


def split(df):
    return {k: df.loc[a:b] for k, (a, b) in SPLIT.items()}


def _stars(p):
    if pd.isna(p):
        return "n/a"
    return f"{p:.4f} " + ("***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "." if p < 0.1 else "ns")


def inference_table(fit, metrics):
    t = fit.params.copy()
    t.insert(0, "parameter/order", fit.spec)
    t.insert(0, "model_name", fit.name)
    t["p_value_or_significance"] = t.pop("p_value").map(_stars)
    t["AIC"], t["BIC"] = fit.aic, fit.bic
    for k, v in metrics.items():
        t[k] = v
    t["notes"] = ""
    t.loc[t.index[0], "notes"] = fit.notes
    return t[INF_COLS]


def _long_fc(fit, actual, split_name):
    f = fit.forecast.copy()
    f["actual"] = actual.reindex(f.index)
    f["model"] = fit.name
    f["split"] = split_name
    return f.rename_axis("date").reset_index()


def run_phase3():
    df = load_daily(model_window=True)
    s = split(df)
    print("split sizes:", {k: len(v) for k, v in s.items()})
    trainval = pd.concat([s["train"], s["val"]])

    fits, tables, diags, val_long, test_long, comp = {}, {}, {}, [], [], []
    for fn in mdl.MODELS:
        t0 = time.time()
        f = fn(s["train"], s["val"])
        m = accuracy(s["val"][TARGET].values, f.forecast["mean"].values)
        d = residual_diagnostics(f.resid, f.n_arma_params)
        img = residual_plot(f.name, f.resid, f.scale)
        ft = fn(trainval, s["test"])                      # sealed test forecast for Phase 4
        fits[f.name], tables[f.name], diags[f.name] = (f, img, ft), inference_table(f, m), d
        val_long.append(_long_fc(f, s["val"][TARGET], "val"))
        test_long.append(_long_fc(ft, s["test"][TARGET], "test"))
        lb14 = d.loc[d.test == "Ljung-Box lag 14", "p_value"].iloc[0]
        jb = d.loc[d.test == "Jarque-Bera normality", "p_value"].iloc[0]
        cover = float(((s["val"][TARGET] >= f.forecast.lower) & (s["val"][TARGET] <= f.forecast.upper)).mean() * 100)
        comp.append(dict(model_name=f.name, specification=f.spec, fit_scale=f.scale, **{f"val_{k}": v for k, v in m.items()},
                         val_95pct_interval_coverage=cover, AIC=f.aic, BIC=f.bic, n_params=len(f.params),
                         ljung_box_lag14_p=lb14, jarque_bera_p=jb,
                         uses_realised_lead_time="yes (oracle)" if "with_lead" in f.name else "no",
                         december_handling="target adjusted by surge multiplier" if fn in mdl.UNIVARIATE else "dec_surge_2023 dummy regressor",
                         test_spec_after_refit=ft.spec))
        print(f"  {f.name:28s} val MAPE {m['MAPE']:6.1f}%  RMSE {m['RMSE']:9,.0f}  [{time.time() - t0:.1f}s]  {f.spec}")

    sens = sensitivity_no_adjust(s)
    cs = pd.DataFrame(comp)
    for k in ("MAPE", "RMSE", "MAE"):
        cs[f"rank_{k}"] = cs[f"val_{k}"].rank(method="min").astype(int)
    cs["mean_rank"] = cs[["rank_MAPE", "rank_RMSE", "rank_MAE"]].mean(axis=1)
    cs = cs.sort_values(["rank_MAPE", "mean_rank"]).reset_index(drop=True)
    cs.insert(0, "overall_rank_by_val_MAPE", range(1, len(cs) + 1))

    pd.concat(val_long).to_csv(OUT_D / "phase3_forecasts_val.csv", index=False)
    pd.concat(test_long).to_csv(OUT_D / "phase3_forecasts_test.csv", index=False)
    _write_excel(fits, tables, diags, cs, s, sens)
    _plot_val(df, s, fits, cs)
    print("Wrote", OUT_X / "model_inference_tables.xlsx")
    return cs, tables, diags, fits


def sensitivity_no_adjust(s):
    """Validation accuracy of the univariate models WITHOUT the December adjustment (kept for transparency)."""
    rows = []
    mdl.ADJUST_DEC_FOR_UNIVARIATE = False
    try:
        for fn in mdl.UNIVARIATE:
            f = fn(s["train"], s["val"])
            rows.append(dict(model_name=f.name, specification=f.spec,
                             **{f"val_{k}_unadjusted": v for k, v in accuracy(s["val"][TARGET].values, f.forecast["mean"].values).items()}))
    finally:
        mdl.ADJUST_DEC_FOR_UNIVARIATE = True
    mult, p = mdl.dec_multiplier(s["train"])
    out = pd.DataFrame(rows)
    out["note"] = ""
    out.loc[0, "note"] = (f"Train ends inside the Dec-2023 surge; unadjusted univariate models carry the surge level into Jan 2024. "
                          f"Adjustment multiplier estimated on train = {mult:.2f}x (p={p:.1g}).")
    return out


def _write_excel(fits, tables, diags, cs, s, sens):
    path = OUT_X / "model_inference_tables.xlsx"
    bold = Font(bold=True)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        cs.to_excel(w, sheet_name="model_comparison_summary", index=False)
        sens.merge(cs[["model_name", "val_MAPE", "val_RMSE", "val_MAE"]].rename(columns=lambda c: c + "_adjusted" if c.startswith("val_") else c),
                   on="model_name").to_excel(w, sheet_name="sensitivity_no_dec_adjust", index=False)
        pd.DataFrame([dict(split=k, start=a, end=b, days=len(s[k]),
                           role={"train": "estimation", "val": "ranking / model selection (multi-step from end of train)",
                                 "test": "sealed - refit on train+val, forecasts saved for Phase 4 only"}[k])
                      for k, (a, b) in SPLIT.items()]).to_excel(w, sheet_name="split_definition", index=False)
        pd.concat(tables.values(), ignore_index=True).to_excel(w, sheet_name="all_inference_tables", index=False)
        for name, t in tables.items():
            sh = name[:31]
            t.to_excel(w, sheet_name=sh, index=False, startrow=1)
            r0 = len(t) + 4
            diags[name].to_excel(w, sheet_name=sh, index=False, startrow=r0 + 1)
            ws = w.sheets[sh]
            ws.cell(row=1, column=1, value=f"{name} — inference table (validation metrics: Jan 2024, multi-step)").font = bold
            ws.cell(row=r0 + 1, column=1, value="Residual diagnostics (in-sample, train)").font = bold
            f = fits[name][0]
            r1 = r0 + len(diags[name]) + 4
            sel = f.extra.get("selection")
            if sel is not None:
                ws.cell(row=r1, column=1, value="Order selection candidates (top 10 by AIC)").font = bold
                sel.astype(str).to_excel(w, sheet_name=sh, index=False, startrow=r1)
                r1 += len(sel) + 3
            ws.cell(row=r1, column=1, value="Residual plot, ACF and Q-Q").font = bold
            ws.add_image(XLImage(str(fits[name][1])), f"A{r1 + 1}")
            for col, wdt in zip("ABCDEFGHIJKL", (26, 44, 30, 12, 12, 16, 11, 11, 11, 11, 9, 80)):
                ws.column_dimensions[col].width = wdt


def _plot_val(df, s, fits, cs):
    order = list(cs.model_name)
    n = len(order)
    cols = 3
    rows = int(np.ceil(n / cols))
    fig, axes = plt.subplots(rows, cols, figsize=(14, 2.6 * rows), sharex=True, sharey=True)
    hist = df.loc["2023-12-01":"2023-12-31", TARGET]
    for ax, name in zip(axes.flat, order):
        f = fits[name][0].forecast
        ax.plot(hist.index, hist, color=C["muted"], lw=1)
        ax.fill_between(f.index, f.lower, f.upper, color=C["s2"], alpha=0.15, lw=0)
        ax.plot(s["val"].index, s["val"][TARGET], color=C["ink"], lw=1.2, label="actual")
        ax.plot(f.index, f["mean"], color=C["s2"], lw=1.6, label="forecast")
        r = cs.set_index("model_name").loc[name]
        ax.set_title(f"#{int(r.overall_rank_by_val_MAPE)} {name} — MAPE {r.val_MAPE:.1f}%", fontsize=9)
        ax.set_ylim(0, 1.3 * max(s["val"][TARGET].max(), 80000))
        _thousands(ax)
    for ax in axes.flat[n:]:
        ax.axis("off")
    axes.flat[0].legend(loc="upper right")
    for ax in axes[-1]:
        _fmt_dates(ax)
    fig.suptitle("Validation (Jan 2024) multi-step forecasts from end of train; grey = Dec 2023 history, band = 95% interval",
                 x=0.01, ha="left", fontsize=11, fontweight="bold")
    fig.tight_layout()
    fig.savefig(OUT_FIG / "phase3_01_val_forecasts.png", dpi=140, bbox_inches="tight")
    plt.close(fig)


if __name__ == "__main__":
    run_phase3()
